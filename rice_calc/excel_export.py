"""計算結果をA4用紙1枚・大きな文字で印刷できるExcelファイルにまとめる。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins

HEADER_FONT = Font(size=28, bold=True)
SECTION_FONT = Font(size=22, bold=True)
LABEL_FONT = Font(size=20)
VALUE_FONT = Font(size=20, bold=True)
TOTAL_FONT = Font(size=22, bold=True)

_THIN = Side(style="thin")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def _write_row(ws, row: int, label: str, value, label_font: Font, value_font: Font, height: int = 30, fmt: str = "0.0"):
    c1 = ws.cell(row=row, column=1, value=label)
    c1.font = label_font
    c1.alignment = LEFT
    c1.border = BORDER
    c2 = ws.cell(row=row, column=2, value=value)
    c2.font = value_font
    c2.alignment = RIGHT
    c2.border = BORDER
    c2.number_format = fmt
    ws.row_dimensions[row].height = height


def build_print_workbook(day_label: str, meal_results: list[dict], soup_water_l: float | None) -> bytes:
    """meal_results: [{"meal": str, "results": {class_name: kg}, "total": float,
    "raw_rice": float | None, "water": float | None}, ...]"""
    wb = Workbook()
    ws = wb.active
    ws.title = "配缶量"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14

    row = 1
    ws.cell(row=row, column=1, value=day_label).font = HEADER_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.cell(row=row, column=1).alignment = CENTER
    ws.row_dimensions[row].height = 42
    row += 2

    for meal in meal_results:
        ws.cell(row=row, column=1, value=meal["meal"]).font = SECTION_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 36
        row += 1

        for class_name, kg in meal["results"].items():
            _write_row(ws, row, class_name, kg, LABEL_FONT, VALUE_FONT)
            row += 1

        _write_row(ws, row, "合計", meal["total"], TOTAL_FONT, TOTAL_FONT, height=34)
        row += 1

        if meal.get("raw_rice") is not None:
            _write_row(ws, row, "生米(kg)", meal["raw_rice"], LABEL_FONT, VALUE_FONT)
            row += 1
            _write_row(ws, row, "水(L)", meal["water"], LABEL_FONT, VALUE_FONT)
            row += 1

        row += 1  # 空行

    if soup_water_l is not None:
        _write_row(ws, row, "味噌汁の水(L)", soup_water_l, SECTION_FONT, TOTAL_FONT, height=36)
        row += 1

    last_row = row - 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    ws.print_area = f"A1:B{last_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
